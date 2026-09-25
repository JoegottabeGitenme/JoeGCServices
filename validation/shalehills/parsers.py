"""Parsers for the real Shale Hills Critical Zone Observatory (SSHCZO)
data, acquired Session 10 for the paper's second validation site
(geowatch.pdf Section 4.2.2: 74 dates, RMSE 0.060 -> 0.054, 55/74 dates
improved -- paper's own citation: Naithani, K., Baldwin, D., 2015. "CZO
Dataset: Shale Hills - Soil Moisture, Hydropedologic Properties
2006-2015").

**How the data was found**: the paper's own cited URL
(criticalzone.org/shale-hills/data/dataset/3001/) is defunct (the CZO
program was succeeded by CZ Net in Dec 2020). The archived listing
(czo-archive.criticalzone.org/shale-hills/data/dataset/3001/) points to
"the updated resource on HydroShare.org" -- a live, unrestricted, CC-BY
licensed resource (id 83efb5317d284ba996ddaeb2b74b6a42), fetched directly
via HydroShare's own REST API (no WAF, unlike Tarrawarra's host). The
Shale Hills DEM, SSURGO soil survey extract, and flux-tower meteorology
were similarly located via HydroShare's public "CZO Shale Hills" group
(group id 147) rather than the original (dead) paper citation.

**A real, load-bearing coordinate-system discrepancy, found and resolved
before writing any interpolation code (Session 10)** -- worth stating up
front since it would otherwise cause a silent, catastrophic misalignment
at this site's scale (an 8ha / ~500x300m catchment):

- The TDR xlsx's own ReadMe.md claims "NAD 1927 State Plane (PA) coordinate
  system." This is WRONG. Decisively confirmed by computing the full
  106-site coordinate cloud's bounding box under two hypotheses (NAD83
  UTM Zone 18N vs. NAD27 UTM Zone 18N) and comparing against HydroShare's
  OWN independently-stated WGS84 lat/lon bbox for this resource: the NAD83
  UTM 18N interpretation matches to 4 decimal degrees EXACTLY; NAD27 is
  off by ~0.002 degrees (~200m north-south -- more than half the
  catchment's own extent). The TDR site coordinates are NAD83 UTM Zone 18N
  (EPSG:26918), full stop -- not State Plane, matching SSURGO's own
  correctly-stated CRS.
- The DEM's own prj.adf (read directly, not assumed) says NAD27 UTM Zone
  18N (EPSG:26718) -- a DIFFERENT datum from the TDR/SSURGO data. The
  NAD27->NAD83 shift at this specific location is (+32m East, +212m
  North) -- confirmed by direct pyproj transformation, not estimated.
  Left uncorrected, every TDR site would silently sample the DEM ~212m
  away from its true location -- more than half the catchment's height.
- **Resolution**: the committed DEM
  (`shalehills_dem_3m_nad83utm18n.tif`) has ALREADY been reprojected to
  EPSG:26918 (NAD83 UTM 18N) via rasterio's warp/reproject
  (bilinear resampling), matching the TDR and SSURGO data's real CRS. This
  is the single most important fact for anyone touching this directory --
  never reintroduce the original NAD27 DEM without reprojecting it first.

**Units, a second real discrepancy worth flagging (opposite direction
from Tarrawarra's own)**: the TDR xlsx's own values are ALREADY fractional
m3/m3 (e.g. 0.2215), matching the paper's own published RMSE units
directly. Tarrawarra's TDR files were %V/V (needing /100 before use) --
reflexively applying that same conversion here would silently divide
every value by 100 a second time. No conversion is applied in
`parse_tdr_xlsx` below; this is deliberate, not an oversight.

**Depth choice**: unlike Tarrawarra's TDR (one integrated ~30cm-average
reading per point) and NMM (discrete but closely-spaced 15/30/45/60/90cm),
Shale Hills' TRIME-T3 tube probes read at 10/20/40/60/80/100cm -- widely
spaced discrete depths with no natural "shallow pair" to average the way
NMM's 15+30cm did. The shallowest depth, 10cm, is used as this project's
primary comparison quantity -- the closest available physical match to
the near-surface value GeoWATCH's Eq. 1 downscales (see
`services/trail-physics/physics/redistribution.py`'s module docstring:
the real production equation's `theta_SMAP` input is a near-surface
satellite/land-surface-model value, not a full-profile average).

**Soil texture, dominant-component convention**: SSURGO map units here are
frequently associations/complexes of 2 named soil series (e.g. "BMF" =
Berks-Weikert association: Berks 50%, Weikert 30%, unnamed remainder).
Following the same "one real, defensible choice, documented, not silently
picked" discipline as Tarrawarra's particle.dat "shallowest layer"
convention: the DOMINANT component (highest `comppct_r`) per map unit is
used, and that component's SHALLOWEST horizon's sand%/clay% feeds the
existing `physics.soil_texture.soil_hydraulic_properties` USDA-triangle ->
Noah SOILPARM.TBL pipeline -- the SAME pipeline Tarrawarra/NMM used,
deliberately, so this remains a genuine test of the frozen Session 8
configuration rather than a different soil-parameter methodology that
happens to also produce a pass.

**SSURGO tabular schema**: the real SSURGO tabular export
(`soildb_US_2002.zip`, a fixed nationwide MS Access template, no per-file
headers) was NOT guessed from memory or an assumed field order --  it was
extracted directly from the actual bundled `soildb_US_2002.mdb` template
database (via the pure-Python `access_parser` package) and independently
verified against a real data row via a checksum property (sand% + silt% +
claytotal% must sum to ~100 for a real horizon; confirmed exactly: 26.3 +
52.7 + 21.0 = 100.0). Only the fields actually used are exposed here;
`SSURGO_MAPUNIT_COLUMNS`/`SSURGO_COMPONENT_COLUMNS`/`SSURGO_CHORIZON_COLUMNS`
below list the FULL real column order for provenance, in case a future
session needs a field not currently parsed.

**Meteorology**: real Campbell Scientific TOA5-format flux-tower logger
files (10-minute intervals, 2009-04-01 through 2013-05-20 -- only 33 of
the 76 TDR dates fall in this window; the rest fall back to Stage-1-only,
exactly like Tarrawarra/NMM's own per-date graceful fallback). Real,
directly-measured net radiation and relative humidity (no Rn-from-Rs
fallback needed, unlike Tarrawarra) but NO wind speed column at all in
this dataset. FAO-56 (Allen et al. 1998, Chapter 3, "Estimating missing
climatic data") explicitly sanctions a fixed default for exactly this
situation: "Where no wind data are available within the region, a value
of 2 m/s can be used as a temporary estimate. This value is the average
over 2000 weather stations around the globe." -- verified against the
live primary source before use, not recalled from memory. See
`DEFAULT_WIND_SPEED_M_S` below.

Real, physically-implausible sensor dropouts are present in the raw
logger files (e.g. a 2011 record with `pressure_irga_mean=11.004` kPa --
obviously a sensor fault, not a real atmospheric pressure at ~300m
elevation). `parse_meteo_file` rejects records with pressure outside a
generous plausible range rather than silently feeding sensor noise into
the FAO-56 calculation.
"""

from __future__ import annotations

import datetime
import re
from dataclasses import dataclass
from pathlib import Path

import numpy as np
import rasterio
import shapefile
from shapely.geometry import Point, shape

# =============================================================================
# TDR soil moisture (SSHCZO_2010TDRSoilMoisture.xlsx)
# =============================================================================

TDR_VALID_DEPTHS_CM = (10, 20, 40, 60, 80, 100)


@dataclass
class TdrSiteReading:
    site_id: str
    x: float  # NAD83 UTM Zone 18N (EPSG:26918) meters -- see module docstring
    y: float
    date: datetime.date
    moisture_frac: float  # m3/m3, NOT %V/V -- see module docstring


def _parse_date_column_header(col_name: str) -> datetime.date:
    """Column headers are 'SM' + a run-together M(M)DDYY date, e.g.
    'SM42510' (4-25-10) or 'SM101010' (10-10-10) -- single-digit months
    give a 5-digit numeric part, two-digit months (Oct/Nov/Dec) give 6.
    Confirmed against the real file: the resulting min/max dates match
    the ReadMe's own stated 'Date Start 2006-12-10' / 'Date End
    2015-07-16' EXACTLY."""
    m = re.match(r"^SM(\d{5,6})$", col_name)
    if not m:
        raise ValueError(f"Unrecognized TDR date column header: {col_name!r}")
    digits = m.group(1)
    if len(digits) == 5:
        month, day, year = int(digits[0]), int(digits[1:3]), int(digits[3:5])
    else:
        month, day, year = int(digits[0:2]), int(digits[2:4]), int(digits[4:6])
    return datetime.date(2000 + year, month, day)


def parse_tdr_xlsx(path: str, depth_cm: int = 10) -> list[TdrSiteReading]:
    """Parse one depth sheet ('SM{depth_cm}cm') of the real TDR workbook.
    Values of 'NA' (a real sensor/probe-refusal gap, common at this site
    given its famously shallow and spatially variable regolith-to-bedrock
    depth) are skipped, not floored to 0 or interpolated.

    **Real data bug found (Session 10), same discipline as Session 4's
    ksat.dat zero-conductivity exclusion and Session 9's NMM negative-
    reading exclusion**: exactly 8 of 4141 readings at 10cm depth are a
    literal 0.0 -- physically implausible for real volumetric soil
    moisture (even bone-dry soil retains residual moisture, typically
    >=0.02-0.05 m3/m3) and confirmed as a data-quality artifact, not a
    real physical measurement, by a clean discontinuity in the value
    distribution: the 8 zeros are followed immediately by a jump to 0.005
    and up -- no smooth continuum near zero the way a real physical
    reading approaching a dry limit would show. All 8 are concentrated at
    just 2 sites (A3, 74B) on specific dates, consistent with a recorded
    probe-fault/refusal code rather than 8 independent real dry readings.
    Excluded here (treated like 'NA'), with a loud, counted warning, not
    silently kept (which would corrupt that date's site-mean baseline and
    RMSE) or silently floored to some other value.
    """
    import openpyxl

    if depth_cm not in TDR_VALID_DEPTHS_CM:
        raise ValueError(f"depth_cm must be one of {TDR_VALID_DEPTHS_CM}, got {depth_cm}")

    wb = openpyxl.load_workbook(path, data_only=True)
    sheet_name = f"SM{depth_cm}cm"
    if sheet_name not in wb.sheetnames:
        raise ValueError(f"Sheet {sheet_name!r} not found in {path} -- have {wb.sheetnames}")
    ws = wb[sheet_name]

    rows = list(ws.iter_rows(values_only=True))
    header = rows[0]
    date_cols = [(i, _parse_date_column_header(c)) for i, c in enumerate(header) if i >= 3]

    readings: list[TdrSiteReading] = []
    n_excluded_zero = 0
    for row in rows[1:]:
        if row[0] is None:
            continue
        site_id, x, y = str(row[0]), float(row[1]), float(row[2])
        for col_idx, date in date_cols:
            value = row[col_idx]
            if value is None or value == "NA" or not isinstance(value, (int, float)):
                continue
            if float(value) == 0.0:
                n_excluded_zero += 1
                continue
            readings.append(TdrSiteReading(site_id=site_id, x=x, y=y, date=date, moisture_frac=float(value)))
    if n_excluded_zero:
        import sys

        print(
            f"NOTE: excluded {n_excluded_zero} literal-0.0 reading(s) at {depth_cm}cm depth "
            f"(a real data-quality artifact, not a physical measurement -- see "
            f"parse_tdr_xlsx's docstring).",
            file=sys.stderr,
        )
    return readings


# =============================================================================
# DEM (already reprojected to EPSG:26918 -- see module docstring)
# =============================================================================


@dataclass
class DemGrid:
    elevation: np.ndarray  # row 0 = north edge, same convention as validation/tarrawarra
    cellsize: float
    xllcorner: float
    yllcorner: float


def parse_dem_geotiff(path: str) -> DemGrid:
    """Read the reprojected Shale Hills DEM. Confirms cellsize is square
    (raises if not -- the rest of this codebase's grid-construction math
    assumes it) and that the CRS is the expected EPSG:26918 (raises with a
    loud, specific message otherwise -- this is exactly the kind of
    silent-misalignment risk documented at length in the module
    docstring, so a wrong/re-reprojected file must fail loudly, not
    silently produce a subtly-shifted grid)."""
    with rasterio.open(path) as src:
        if src.crs is None or src.crs.to_epsg() != 26918:
            raise ValueError(
                f"{path}: expected CRS EPSG:26918 (NAD83 UTM Zone 18N -- the TDR/SSURGO "
                f"data's real CRS, see module docstring), got {src.crs}. Do not use the "
                f"original NAD27 DEM without reprojecting it first."
            )
        cellsize_x, cellsize_y = src.res
        if abs(cellsize_x - cellsize_y) > 1e-6:
            raise ValueError(f"{path}: non-square cell size {src.res} not supported")
        elevation = src.read(1)
        nodata = src.nodata
        if nodata is not None:
            elevation = np.where(elevation == nodata, np.nan, elevation)
        xllcorner = src.bounds.left
        yllcorner = src.bounds.bottom
    return DemGrid(elevation=elevation, cellsize=float(cellsize_x), xllcorner=xllcorner, yllcorner=yllcorner)


# =============================================================================
# SSURGO soil survey extract (real classic SSURGO tabular + shapefile
# export, subsetted to the ~10 map units intersecting the Shale Hills
# catchment -- see module docstring for schema provenance).
# =============================================================================

# Full real column order for mapunit.txt, extracted directly from the
# bundled soildb_US_2002.mdb template (Session 10) -- only a few of these
# are actually used below; kept complete for provenance / future use.
SSURGO_MAPUNIT_COLUMNS = [
    "musym", "muname", "mukind", "mustatus", "muacres", "mapunitlfw_l", "mapunitlfw_r",
    "mapunitlfw_h", "mapunitpfa_l", "mapunitpfa_r", "mapunitpfa_h", "farmlndcl", "muhelcl",
    "muwathelcl", "muwndhelcl", "interpfocus", "invesintens", "iacornsr", "nhiforsoigrp",
    "nhspiagr", "vtsepticsyscl", "mucertstat", "lkey", "mukey",
]

SSURGO_COMPONENT_COLUMNS = [
    "comppct_l", "comppct_r", "comppct_h", "compname", "compkind", "majcompflag", "otherph",
    "localphase", "slope_l", "slope_r", "slope_h", "slopelenusle_l", "slopelenusle_r",
    "slopelenusle_h", "runoff", "tfact", "wei", "weg", "erocl", "earthcovkind1", "earthcovkind2",
    "hydricon", "hydricrating", "drainagecl", "elev_l", "elev_r", "elev_h", "aspectccwise",
    "aspectrep", "aspectcwise", "geomdesc", "albedodry_l", "albedodry_r", "albedodry_h",
    "airtempa_l", "airtempa_r", "airtempa_h", "map_l", "map_r", "map_h", "reannualprecip_l",
    "reannualprecip_r", "reannualprecip_h", "ffd_l", "ffd_r", "ffd_h", "nirrcapcl", "nirrcapscl",
    "nirrcapunit", "irrcapcl", "irrcapscl", "irrcapunit", "cropprodindex", "constreeshrubgrp",
    "wndbrksuitgrp", "rsprod_l", "rsprod_r", "rsprod_h", "foragesuitgrpid", "wlgrain", "wlgrass",
    "wlherbaceous", "wlshrub", "wlconiferous", "wlhardwood", "wlwetplant", "wlshallowwat",
    "wlrangeland", "wlopenland", "wlwoodland", "wlwetland", "soilslippot", "frostact",
    "initsub_l", "initsub_r", "initsub_h", "totalsub_l", "totalsub_r", "totalsub_h", "hydgrp",
    "corcon", "corsteel", "taxclname", "taxorder", "taxsuborder", "taxgrtgroup", "taxsubgrp",
    "taxpartsize", "taxpartsizemod", "taxceactcl", "taxreaction", "taxtempcl", "taxmoistscl",
    "taxtempregime", "soiltaxedition", "castorieindex", "flecolcomnum", "flhe", "flphe",
    "flsoilleachpot", "flsoirunoffpot", "fltemik2use", "fltriumph2use", "indraingrp",
    "innitrateleachi", "misoimgmtgrp", "vasoimgtgrp", "mukey", "cokey",
]

SSURGO_CHORIZON_COLUMNS = [
    "hzname", "desgndisc", "desgnmaster", "desgnmasterprime", "desgnvert", "hzdept_l",
    "hzdept_r", "hzdept_h", "hzdepb_l", "hzdepb_r", "hzdepb_h", "hzthk_l", "hzthk_r", "hzthk_h",
    "fraggt10_l", "fraggt10_r", "fraggt10_h", "frag3to10_l", "frag3to10_r", "frag3to10_h",
    "sieveno4_l", "sieveno4_r", "sieveno4_h", "sieveno10_l", "sieveno10_r", "sieveno10_h",
    "sieveno40_l", "sieveno40_r", "sieveno40_h", "sieveno200_l", "sieveno200_r", "sieveno200_h",
    "sandtotal_l", "sandtotal_r", "sandtotal_h", "sandvc_l", "sandvc_r", "sandvc_h", "sandco_l",
    "sandco_r", "sandco_h", "sandmed_l", "sandmed_r", "sandmed_h", "sandfine_l", "sandfine_r",
    "sandfine_h", "sandvf_l", "sandvf_r", "sandvf_h", "silttotal_l", "silttotal_r", "silttotal_h",
    "siltco_l", "siltco_r", "siltco_h", "siltfine_l", "siltfine_r", "siltfine_h", "claytotal_l",
    "claytotal_r", "claytotal_h", "claysizedcarb_l", "claysizedcarb_r", "claysizedcarb_h",
    "om_l", "om_r", "om_h", "dbtenthbar_l", "dbtenthbar_r", "dbtenthbar_h", "dbthirdbar_l",
    "dbthirdbar_r", "dbthirdbar_h", "dbfifteenbar_l", "dbfifteenbar_r", "dbfifteenbar_h",
    "dbovendry_l", "dbovendry_r", "dbovendry_h", "partdensity", "ksat_l", "ksat_r", "ksat_h",
    "awc_l", "awc_r", "awc_h", "wtenthbar_l", "wtenthbar_r", "wtenthbar_h", "wthirdbar_l",
    "wthirdbar_r", "wthirdbar_h", "wfifteenbar_l", "wfifteenbar_r", "wfifteenbar_h",
    "wsatiated_l", "wsatiated_r", "wsatiated_h", "lep_l", "lep_r", "lep_h", "ll_l", "ll_r",
    "ll_h", "pi_l", "pi_r", "pi_h", "aashind_l", "aashind_r", "aashind_h", "kwfact", "kffact",
    "caco3_l", "caco3_r", "caco3_h", "gypsum_l", "gypsum_r", "gypsum_h", "sar_l", "sar_r",
    "sar_h", "ec_l", "ec_r", "ec_h", "cec7_l", "cec7_r", "cec7_h", "ecec_l", "ecec_r", "ecec_h",
    "sumbases_l", "sumbases_r", "sumbases_h", "ph1to1h2o_l", "ph1to1h2o_r", "ph1to1h2o_h",
    "ph01mcacl2_l", "ph01mcacl2_r", "ph01mcacl2_h", "freeiron_l", "freeiron_r", "freeiron_h",
    "feoxalate_l", "feoxalate_r", "feoxalate_h", "extracid_l", "extracid_r", "extracid_h",
    "extral_l", "extral_r", "extral_h", "aloxalate_l", "aloxalate_r", "aloxalate_h", "pbray1_l",
    "pbray1_r", "pbray1_h", "poxalate_l", "poxalate_r", "poxalate_h", "ph2osoluble_l",
    "ph2osoluble_r", "ph2osoluble_h", "ptotal_l", "ptotal_r", "ptotal_h", "excavdifcl",
    "excavdifms", "cokey", "chkey",
]


def _parse_pipe_delimited(path: str, columns: list[str]) -> list[dict]:
    """Real classic-SSURGO tabular export: pipe-delimited, double-quoted
    text fields, no header row, latin-1 encoded (confirmed against the
    real files -- some muname/compname fields contain non-ASCII
    characters)."""
    rows = []
    with open(path, encoding="latin-1") as f:
        for line in f:
            fields = [x.strip('"') for x in line.rstrip("\n").split("|")]
            rows.append(dict(zip(columns, fields)))
    return rows


def parse_ssurgo_mapunit(path: str) -> list[dict]:
    return _parse_pipe_delimited(path, SSURGO_MAPUNIT_COLUMNS)


def parse_ssurgo_component(path: str) -> list[dict]:
    return _parse_pipe_delimited(path, SSURGO_COMPONENT_COLUMNS)


def parse_ssurgo_chorizon(path: str) -> list[dict]:
    return _parse_pipe_delimited(path, SSURGO_CHORIZON_COLUMNS)


def parse_ssurgo_mapunit_polygons(shp_path: str) -> list[tuple[str, object]]:
    """Returns (mukey, shapely_polygon) for every polygon in the (already
    catchment-subsetted, see README.md) map unit shapefile. CRS is NAD83
    UTM Zone 18N (EPSG:26918) per this shapefile's own .prj and the SSURGO
    readme.txt's own stated 'Coordinate System: UTM Zone 18, Northern
    Hemisphere (NAD 83)' -- consistent with the TDR data and the
    reprojected DEM, unlike the DEM's ORIGINAL (NAD27) file."""
    sf = shapefile.Reader(shp_path)
    result = []
    for sr in sf.iterShapeRecords():
        mukey = sr.record["mukey"]
        geom = shape(sr.shape.__geo_interface__)
        result.append((mukey, geom))
    return result


def dominant_component_texture_by_mukey(
    mapunit_path: str, component_path: str, chorizon_path: str
) -> dict[str, tuple[float, float]]:
    """For each map unit, find its DOMINANT component (highest
    `comppct_r`, the standard SSURGO convention for a single representative
    soil -- some real map units here are named associations/complexes of
    2 series, e.g. 'Berks-Weikert association', see module docstring),
    then that component's SHALLOWEST horizon's (sand%, clay%) -- the same
    "shallowest surface layer" convention Tarrawarra's particle.dat parsing
    established.

    Returns {mukey: (sand_pct, clay_pct)}.
    """
    components = parse_ssurgo_component(component_path)
    horizons = parse_ssurgo_chorizon(chorizon_path)
    horizons_by_cokey: dict[str, list[dict]] = {}
    for h in horizons:
        horizons_by_cokey.setdefault(h["cokey"], []).append(h)

    dominant_cokey_by_mukey: dict[str, str] = {}
    best_pct_by_mukey: dict[str, float] = {}
    for c in components:
        mukey = c["mukey"]
        pct = float(c["comppct_r"]) if c["comppct_r"] else 0.0
        if mukey not in best_pct_by_mukey or pct > best_pct_by_mukey[mukey]:
            best_pct_by_mukey[mukey] = pct
            dominant_cokey_by_mukey[mukey] = c["cokey"]

    result: dict[str, tuple[float, float]] = {}
    for mukey, cokey in dominant_cokey_by_mukey.items():
        layers = horizons_by_cokey.get(cokey, [])
        if not layers:
            continue
        layers.sort(key=lambda h: float(h["hzdept_r"]) if h["hzdept_r"] else 0.0)
        shallowest = layers[0]
        if not shallowest["sandtotal_r"] or not shallowest["claytotal_r"]:
            continue
        result[mukey] = (float(shallowest["sandtotal_r"]), float(shallowest["claytotal_r"]))
    return result


def assign_mukey_to_points(
    polygons: list[tuple[str, object]], points_xy: np.ndarray
) -> list[str | None]:
    """Nearest-polygon assignment: for each point, the mukey of the
    polygon it falls inside, or (if it falls in none -- possible right at
    the catchment/DEM edge, since the map-unit polygons and the DEM extent
    don't share an exact boundary) the mukey of the NEAREST polygon,
    loudly counted, rather than silently dropping the point."""
    mukeys: list[str | None] = []
    n_fallback = 0
    for x, y in points_xy:
        pt = Point(x, y)
        found = None
        for mukey, geom in polygons:
            if geom.contains(pt):
                found = mukey
                break
        if found is None:
            n_fallback += 1
            best_mukey, best_dist = None, float("inf")
            for mukey, geom in polygons:
                d = geom.distance(pt)
                if d < best_dist:
                    best_dist, best_mukey = d, mukey
            found = best_mukey
        mukeys.append(found)
    if n_fallback:
        import sys

        print(
            f"NOTE: {n_fallback}/{len(points_xy)} point(s) did not fall inside any map unit "
            f"polygon -- assigned to the nearest one instead (expected near the catchment edge).",
            file=sys.stderr,
        )
    return mukeys


# =============================================================================
# Meteorology (real Campbell Scientific TOA5 flux-tower logger files)
# =============================================================================

# FAO-56 (Allen et al. 1998), Chapter 3, "Estimating missing climatic
# data": "Where no wind data are available within the region, a value of
# 2 m/s can be used as a temporary estimate. This value is the average
# over 2000 weather stations around the globe." Verified against the live
# primary source (fao.org/4/x0490e/x0490e07.htm) Session 10 -- this
# dataset's flux tower has no wind speed sensor at all.
DEFAULT_WIND_SPEED_M_S = 2.0

# Real sensor dropouts exist in the raw logger files (e.g. a 2011 record
# with pressure_irga_mean=11.004 kPa -- physically impossible at this
# site's ~300m elevation, where true atmospheric pressure is ~97-98 kPa).
# A generous plausible range, not a tight one, so real weather variation
# is never rejected -- only genuine sensor faults.
PLAUSIBLE_PRESSURE_KPA = (70.0, 105.0)


@dataclass
class MeteoRecord:
    timestamp: datetime.datetime
    pressure_kpa: float | None
    temperature_c: float | None
    relative_humidity_pct: float | None
    net_radiation_w_m2: float | None


def parse_meteo_file(path: str) -> list[MeteoRecord]:
    """Parse a real 20XX_CZO_FluxTowerMeteo.dat file. Campbell Scientific
    TOA5 format: 4 header lines (station info, column names, units,
    aggregation-type), then 10-minute-interval comma-separated data.
    Columns used: TIMESTAMP, pressure_irga_mean (kPa), T_hmp_mean (C),
    RH_hmp_current (%), net_radiation_mean (W/m2). Records with a
    physically-implausible pressure (see PLAUSIBLE_PRESSURE_KPA) have
    ALL fields set to None (a real sensor fault casts doubt on the whole
    scan, not just the pressure channel) rather than being silently kept
    or dropped outright (dropping would shift which timestamps exist;
    None-ing preserves the record's existence for daily aggregation logic
    to see "this scan had no usable data", same discipline as
    Tarrawarra's daily.met '*' handling).
    """
    records = []
    with open(path) as f:
        lines = f.readlines()
    for line in lines[4:]:  # skip the 4 TOA5 header lines
        parts = [p.strip().strip('"') for p in line.rstrip("\n").split(",")]
        if len(parts) < 11:
            continue
        try:
            ts = datetime.datetime.strptime(parts[0], "%m/%d/%Y %H:%M")
        except ValueError:
            continue

        def field(v: str) -> float | None:
            if v in ("NAN", "NULL", ""):
                return None
            try:
                return float(v)
            except ValueError:
                return None

        pressure = field(parts[2])
        temp = field(parts[6])
        rh = field(parts[8])
        rn = field(parts[9])

        if pressure is not None and not (PLAUSIBLE_PRESSURE_KPA[0] <= pressure <= PLAUSIBLE_PRESSURE_KPA[1]):
            pressure = temp = rh = rn = None

        records.append(
            MeteoRecord(
                timestamp=ts,
                pressure_kpa=pressure,
                temperature_c=temp,
                relative_humidity_pct=rh,
                net_radiation_w_m2=rn,
            )
        )
    return records
